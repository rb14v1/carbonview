from rest_framework import serializers

from rest_framework.views import APIView
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, Reference
from django.db.models import Sum
from calendar import month_name

from apps.company.models import Organization, ReportingPeriod
from apps.emission.models import EmissionActivity, EmissionScope, RecordStatus
from apps.summary.models import EnergySummary, WaterUsage, WasteSummary

from ..models import (
    AuditLog,
    CarbonCredit,
    Certificate,
    EnergySummary,
    MonthlyEmissionSummary,
    NeutralityClaim,
    ReductionProject,
    WasteSummary,
    WaterUsage,
)


class AuditLogSerializer(serializers.ModelSerializer):
    content_type_name = serializers.CharField(source="content_type.model", read_only=True)
    changed_by_name = serializers.CharField(source="changed_by.username", read_only=True)

    class Meta:
        model = AuditLog
        fields = [
            "id",
            "content_type",
            "content_type_name",
            "object_id",
            "action",
            "message",
            "changed_by",
            "changed_by_name",
            "created_at",
            "updated_at",
        ]
        read_only_fields = ["id", "content_type_name", "changed_by_name", "created_at", "updated_at"]


class CertificateSerializer(serializers.ModelSerializer):
    file_url = serializers.SerializerMethodField()

    class Meta:
        model = Certificate
        fields = [
            "id",
            "name",
            "description",
            "file",
            "file_url",
            "created_at",
            "updated_at",
        ]
        read_only_fields = ["id", "file_url", "created_at", "updated_at"]

    def get_file_url(self, obj):
        """Build an absolute file URL when a request is available and fall back to the stored relative path otherwise."""
        request = self.context.get("request")
        if not obj.file:
            return None
        url = obj.file.url
        return request.build_absolute_uri(url) if request else url

    def validate_file(self, value):
        """Allow only supported image uploads and reject files that exceed the configured size limit."""
        allowed_types = {"image/jpeg", "image/jpg", "image/png", "image/webp"}
        content_type = getattr(value, "content_type", None)
        if content_type and content_type not in allowed_types:
            raise serializers.ValidationError("Only JPG, PNG, and WEBP images are allowed.")
        max_size = 5 * 1024 * 1024
        if value.size > max_size:
            raise serializers.ValidationError("Image size must be 5 MB or smaller.")
        return value


class CarbonCreditSerializer(serializers.ModelSerializer):
    facility_name = serializers.CharField(source="facility.name", read_only=True)

    class Meta:
        model = CarbonCredit
        fields = [
            "id",
            "facility",
            "facility_name",
            "project_name",
            "source",
            "registry_name",
            "serial_number",
            "vintage_year",
            "quantity_tco2e",
            "purchase_date",
            "retirement_date",
            "status",
            "created_at",
            "updated_at",
        ]
        read_only_fields = ["id", "facility_name", "created_at", "updated_at"]


class ReductionProjectSerializer(serializers.ModelSerializer):
    facility_name = serializers.CharField(source="facility.name", read_only=True)

    class Meta:
        model = ReductionProject
        fields = [
            "id",
            "facility",
            "facility_name",
            "name",
            "description",
            "start_date",
            "end_date",
            "expected_reduction_tco2e",
            "actual_reduction_tco2e",
            "created_at",
            "updated_at",
        ]
        read_only_fields = ["id", "facility_name", "created_at", "updated_at"]


class NeutralityClaimSerializer(serializers.ModelSerializer):
    organization_name = serializers.CharField(source="organization.name", read_only=True)
    facility_name = serializers.CharField(source="facility.name", read_only=True)
    reporting_period_name = serializers.CharField(source="reporting_period.name", read_only=True)

    class Meta:
        model = NeutralityClaim
        fields = [
            "id",
            "organization",
            "organization_name",
            "facility",
            "facility_name",
            "reporting_period",
            "reporting_period_name",
            "name",
            "gross_emissions",
            "reductions",
            "offsets_retired",
            "net_emissions",
            "status",
            "notes",
            "created_at",
            "updated_at",
        ]
        read_only_fields = [
            "id",
            "organization_name",
            "facility_name",
            "reporting_period_name",
            "gross_emissions",
            "reductions",
            "offsets_retired",
            "net_emissions",
            "created_at",
            "updated_at",
        ]


class MonthlyEmissionSummarySerializer(serializers.ModelSerializer):
    organization_name = serializers.CharField(source="organization.name", read_only=True)
    facility_name = serializers.CharField(source="facility.name", read_only=True)
    reporting_period_name = serializers.CharField(source="reporting_period.name", read_only=True)

    class Meta:
        model = MonthlyEmissionSummary
        fields = [
            "id",
            "organization",
            "organization_name",
            "facility",
            "facility_name",
            "reporting_period",
            "reporting_period_name",
            "year",
            "month",
            "scope",
            "total_co2e",
            "created_at",
            "updated_at",
        ]
        read_only_fields = [
            "id",
            "organization_name",
            "facility_name",
            "reporting_period_name",
            "created_at",
            "updated_at",
        ]


class EnergySummarySerializer(serializers.ModelSerializer):
    facility_name = serializers.CharField(source="facility.name", read_only=True)
    reporting_period_name = serializers.CharField(source="reporting_period.name", read_only=True)
    unit_code = serializers.CharField(source="unit.code", read_only=True)

    class Meta:
        model = EnergySummary
        fields = [
            "id",
            "facility",
            "facility_name",
            "reporting_period",
            "reporting_period_name",
            "total_energy",
            "renewable_energy",
            "non_renewable_energy",
            "unit",
            "unit_code",
            "created_at",
            "updated_at",
        ]
        read_only_fields = [
            "id",
            "facility_name",
            "reporting_period_name",
            "unit_code",
            "created_at",
            "updated_at",
        ]


class WaterUsageSerializer(serializers.ModelSerializer):
    facility_name = serializers.CharField(source="facility.name", read_only=True)
    reporting_period_name = serializers.CharField(source="reporting_period.name", read_only=True)
    unit_code = serializers.CharField(source="unit.code", read_only=True)

    class Meta:
        model = WaterUsage
        fields = [
            "id",
            "facility",
            "facility_name",
            "reporting_period",
            "reporting_period_name",
            "category",
            "amount",
            "unit",
            "unit_code",
            "created_at",
            "updated_at",
        ]
        read_only_fields = [
            "id",
            "facility_name",
            "reporting_period_name",
            "unit_code",
            "created_at",
            "updated_at",
        ]


class WasteSummarySerializer(serializers.ModelSerializer):
    facility_name = serializers.CharField(source="facility.name", read_only=True)
    reporting_period_name = serializers.CharField(source="reporting_period.name", read_only=True)
    unit_code = serializers.CharField(source="unit.code", read_only=True)

    class Meta:
        model = WasteSummary
        fields = [
            "id",
            "facility",
            "facility_name",
            "reporting_period",
            "reporting_period_name",
            "total_waste",
            "hazardous",
            "non_hazardous",
            "recycled_percentage",
            "unit",
            "unit_code",
            "created_at",
            "updated_at",
        ]
        read_only_fields = [
            "id",
            "facility_name",
            "reporting_period_name",
            "unit_code",
            "created_at",
            "updated_at",
        ]



class ESGReportExcelView(APIView):

    def get(self, request):
        """Generate a year-based ESG workbook by aggregating approved emissions plus energy, water, and waste summaries."""
        year = request.GET.get("year")
        org_id = request.GET.get("organization")

        if not year:
            return HttpResponse("Year is required", status=400)

        wb = Workbook()
        ws = wb.active
        ws.title = "ESG Report"

        # =========================
        # Styling helpers
        # =========================
        title_font = Font(size=16, bold=True)
        header_font = Font(bold=True)
        center_align = Alignment(horizontal="center")

        row = 1

        # =========================
        # TITLE
        # =========================
        ws.cell(row=row, column=1, value=f"ESG REPORT - {year}").font = title_font
        row += 2

        # =========================
        # KPI SUMMARY
        # =========================
        ws.cell(row=row, column=1, value="KPI SUMMARY").font = header_font
        row += 1

        reporting_periods = ReportingPeriod.objects.filter(year=year)

        emissions = EmissionActivity.objects.filter(
            reporting_period__in=reporting_periods,
            status=RecordStatus.APPROVED
        )

        if org_id:
            emissions = emissions.filter(organization_id=org_id)

        total_emission = emissions.aggregate(total=Sum("total_co2e"))["total"] or 0

        energy = EnergySummary.objects.filter(reporting_period__in=reporting_periods)
        water = WaterUsage.objects.filter(reporting_period__in=reporting_periods)
        waste = WasteSummary.objects.filter(reporting_period__in=reporting_periods)

        ws.append(["Total Emissions (tCO2e)", float(total_emission)])
        ws.append(["Total Energy", float(energy.aggregate(Sum("total_energy"))["total_energy__sum"] or 0)])
        ws.append(["Total Water", float(water.aggregate(Sum("amount"))["amount__sum"] or 0)])
        ws.append(["Total Waste", float(waste.aggregate(Sum("total_waste"))["total_waste__sum"] or 0)])

        row += 6

        # =========================
        # EMISSIONS SECTION
        # =========================
        ws.cell(row=row, column=1, value="GHG EMISSIONS").font = header_font
        row += 1

        ws.append(["Scope", "Total CO2e"])

        scope_data = (
            emissions.values("category__scope")
            .annotate(total=Sum("total_co2e"))
        )

        scope_row_start = ws.max_row + 1

        for s in scope_data:
            ws.append([s["category__scope"], float(s["total"])])

        scope_row_end = ws.max_row

        row = ws.max_row + 3

        # =========================
        # MONTHLY TREND
        # =========================
        ws.cell(row=row, column=1, value="Monthly Emissions Trend").font = header_font
        row += 1

        ws.append(["Month", "Emissions"])

        monthly_data = (
            emissions.values("activity_date__month")
            .annotate(total=Sum("total_co2e"))
            .order_by("activity_date__month")
        )

        month_row_start = ws.max_row + 1

        for m in monthly_data:
            ws.append([month_name[m["activity_date__month"]], float(m["total"])])

        month_row_end = ws.max_row

        row = ws.max_row + 3

        # =========================
        # ENERGY SECTION
        # =========================
        ws.cell(row=row, column=1, value="ENERGY").font = header_font
        row += 1

        ws.append(["Facility", "Total", "Renewable", "Non-Renewable"])

        for e in energy:
            ws.append([
                str(e.facility),
                float(e.total_energy),
                float(e.renewable_energy),
                float(e.non_renewable_energy),
            ])

        row = ws.max_row + 3

        # =========================
        # WATER SECTION
        # =========================
        ws.cell(row=row, column=1, value="WATER").font = header_font
        row += 1

        ws.append(["Facility", "Category", "Amount"])

        for w in water:
            ws.append([str(w.facility), w.category, float(w.amount)])

        row = ws.max_row + 3

        # =========================
        # WASTE SECTION
        # =========================
        ws.cell(row=row, column=1, value="WASTE").font = header_font
        row += 1

        ws.append(["Facility", "Total", "Hazardous", "Non-Hazardous"])

        for w in waste:
            ws.append([
                str(w.facility),
                float(w.total_waste),
                float(w.hazardous),
                float(w.non_hazardous),
            ])

        # =========================
        # CHART 1: Scope Chart
        # =========================
        chart1 = BarChart()
        chart1.title = "Emissions by Scope"

        data = Reference(ws, min_col=2, min_row=scope_row_start-1, max_row=scope_row_end)
        categories = Reference(ws, min_col=1, min_row=scope_row_start, max_row=scope_row_end)

        chart1.add_data(data, titles_from_data=True)
        chart1.set_categories(categories)

        ws.add_chart(chart1, "H5")

        # =========================
        # CHART 2: Monthly Trend
        # =========================
        chart2 = BarChart()
        chart2.title = "Monthly Emissions Trend"

        data = Reference(ws, min_col=2, min_row=month_row_start-1, max_row=month_row_end)
        categories = Reference(ws, min_col=1, min_row=month_row_start, max_row=month_row_end)

        chart2.add_data(data, titles_from_data=True)
        chart2.set_categories(categories)

        ws.add_chart(chart2, "H20")

        # =========================
        # RESPONSE
        # =========================
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="ESG_Report_{year}.xlsx"'

        wb.save(response)
        return response
