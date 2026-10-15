import calendar
from calendar import month_name
from decimal import Decimal
from io import BytesIO

from django.db.models import Count, Q, Sum
from django.db.models.functions import ExtractMonth, ExtractYear
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework import viewsets
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.views import APIView

from apps.company.models import Facility, Organization, ReportingPeriod
from apps.emission.models import (
    EmissionActivity,
    EmissionCalculation,
    EmissionScope,
    RecordStatus,
)

from .models import (
    AuditLog,
    CarbonCredit,
    Certificate,
    EnergySummary,
    MonthlyEmissionSummary,
    NeutralityClaim,
    ReductionProject,
    WasteSummary,
    WaterUsage,
)
from .serializers import (
    AuditLogSerializer,
    CarbonCreditSerializer,
    CertificateSerializer,
    EnergySummarySerializer,
    MonthlyEmissionSummarySerializer,
    NeutralityClaimSerializer,
    ReductionProjectSerializer,
    WasteSummarySerializer,
    WaterUsageSerializer,
)

class AuditLogViewSet(viewsets.ModelViewSet):
    queryset = AuditLog.objects.select_related("content_type", "changed_by").all()
    serializer_class = AuditLogSerializer


class CertificateViewSet(viewsets.ModelViewSet):
    queryset = Certificate.objects.all()
    serializer_class = CertificateSerializer
    parser_classes = [MultiPartParser, FormParser]


class CarbonCreditViewSet(viewsets.ModelViewSet):
    queryset = CarbonCredit.objects.select_related("facility").all()
    serializer_class = CarbonCreditSerializer


class ReductionProjectViewSet(viewsets.ModelViewSet):
    queryset = ReductionProject.objects.select_related("facility").all()
    serializer_class = ReductionProjectSerializer


class NeutralityClaimViewSet(viewsets.ModelViewSet):
    queryset = NeutralityClaim.objects.select_related(
        "organization", "facility", "reporting_period"
    ).all()
    serializer_class = NeutralityClaimSerializer


class MonthlyEmissionSummaryViewSet(viewsets.ModelViewSet):
    queryset = MonthlyEmissionSummary.objects.select_related(
        "organization", "facility", "reporting_period"
    ).all().order_by("-year")
    serializer_class = MonthlyEmissionSummarySerializer


class EnergySummaryViewSet(viewsets.ModelViewSet):
    queryset = EnergySummary.objects.select_related("facility", "reporting_period", "unit").all()
    serializer_class = EnergySummarySerializer


class WaterUsageViewSet(viewsets.ModelViewSet):
    queryset = WaterUsage.objects.select_related("facility", "reporting_period", "unit").all()
    serializer_class = WaterUsageSerializer


class WasteSummaryViewSet(viewsets.ModelViewSet):
    queryset = WasteSummary.objects.select_related("facility", "reporting_period", "unit").all()
    serializer_class = WasteSummarySerializer

class ESGReportExcelView(APIView):
    def get(self, request):
        """Build a detailed ESG export for the selected date range by aggregating emissions, summaries, claims, and charts."""
        start_date_raw = request.GET.get("start_date")
        end_date_raw = request.GET.get("end_date")
        org_id = request.GET.get("organization")
        facility_id = request.GET.get("facility")

        start_date = parse_date(start_date_raw) if start_date_raw else None
        end_date = parse_date(end_date_raw) if end_date_raw else None

        if not start_date or not end_date:
            return HttpResponse(
                "start_date and end_date are required. Format: YYYY-MM-DD",
                status=400
            )

        if start_date > end_date:
            return HttpResponse(
                "start_date must be earlier than or equal to end_date.",
                status=400
            )

        reporting_periods = ReportingPeriod.objects.filter(
            start_date__lte=end_date,
            end_date__gte=start_date,
        )

        org = None
        facility = None

        if org_id:
            org = get_object_or_404(Organization, id=org_id)
            reporting_periods = reporting_periods.filter(organization=org)

        emissions = EmissionActivity.objects.select_related(
            "organization",
            "facility",
            "department",
            "reporting_period",
            "category",
            "scope3_category",
            "unit",
        ).filter(
            activity_date__range=[start_date, end_date],
            status=RecordStatus.APPROVED,
        )

        if org:
            emissions = emissions.filter(organization=org)

        if facility_id:
            facility = get_object_or_404(Facility, id=facility_id)
            emissions = emissions.filter(facility=facility)

        if org and facility and facility.organization_id != org.id:
            return HttpResponse(
                "Selected facility does not belong to the selected organization",
                status=400
            )

        if facility:
            energy_qs = EnergySummary.objects.filter(
                facility=facility,
                reporting_period__in=reporting_periods
            )
            water_qs = WaterUsage.objects.filter(
                facility=facility,
                reporting_period__in=reporting_periods
            )
            waste_qs = WasteSummary.objects.filter(
                facility=facility,
                reporting_period__in=reporting_periods
            )
            reduction_qs = ReductionProject.objects.filter(
                facility=facility
            ).filter(
                Q(start_date__isnull=True) | Q(start_date__lte=end_date),
                Q(end_date__isnull=True) | Q(end_date__gte=start_date),
            )
            credit_qs = CarbonCredit.objects.filter(
                facility=facility,
                status="retired"
            ).filter(
                Q(retirement_date__isnull=True) | Q(retirement_date__range=[start_date, end_date])
            )
            neutrality_qs = NeutralityClaim.objects.filter(
                facility=facility,
                reporting_period__in=reporting_periods
            )
        elif org:
            energy_qs = EnergySummary.objects.filter(
                facility__organization=org,
                reporting_period__in=reporting_periods
            )
            water_qs = WaterUsage.objects.filter(
                facility__organization=org,
                reporting_period__in=reporting_periods
            )
            waste_qs = WasteSummary.objects.filter(
                facility__organization=org,
                reporting_period__in=reporting_periods
            )
            reduction_qs = ReductionProject.objects.filter(
                facility__organization=org
            ).filter(
                Q(start_date__isnull=True) | Q(start_date__lte=end_date),
                Q(end_date__isnull=True) | Q(end_date__gte=start_date),
            )
            credit_qs = CarbonCredit.objects.filter(
                facility__organization=org,
                status="retired"
            ).filter(
                Q(retirement_date__isnull=True) | Q(retirement_date__range=[start_date, end_date])
            )
            neutrality_qs = NeutralityClaim.objects.filter(
                organization=org,
                reporting_period__in=reporting_periods
            )
        else:
            energy_qs = EnergySummary.objects.filter(reporting_period__in=reporting_periods)
            water_qs = WaterUsage.objects.filter(reporting_period__in=reporting_periods)
            waste_qs = WasteSummary.objects.filter(reporting_period__in=reporting_periods)
            reduction_qs = ReductionProject.objects.filter(
                Q(start_date__isnull=True) | Q(start_date__lte=end_date),
                Q(end_date__isnull=True) | Q(end_date__gte=start_date),
            )
            credit_qs = CarbonCredit.objects.filter(status="retired").filter(
                Q(retirement_date__isnull=True) | Q(retirement_date__range=[start_date, end_date])
            )
            neutrality_qs = NeutralityClaim.objects.filter(reporting_period__in=reporting_periods)

        calculation_qs = EmissionCalculation.objects.select_related(
            "activity__organization",
            "activity__facility",
            "activity__department",
            "activity__reporting_period",
            "activity__category",
            "activity__scope3_category",
            "gas_type",
        ).filter(
            activity__activity_date__range=[start_date, end_date],
            activity__status=RecordStatus.APPROVED,
        )

        if org:
            calculation_qs = calculation_qs.filter(activity__organization=org)
        if facility:
            calculation_qs = calculation_qs.filter(activity__facility=facility)

        total_emissions = calculation_qs.aggregate(total=Sum("co2e_amount"))["total"] or Decimal("0")
        s1_total = calculation_qs.filter(activity__category__scope=EmissionScope.SCOPE_1).aggregate(total=Sum("co2e_amount"))["total"] or Decimal("0")
        s2_total = calculation_qs.filter(activity__category__scope=EmissionScope.SCOPE_2).aggregate(total=Sum("co2e_amount"))["total"] or Decimal("0")
        s3_total = calculation_qs.filter(activity__category__scope=EmissionScope.SCOPE_3).aggregate(total=Sum("co2e_amount"))["total"] or Decimal("0")

        total_energy = energy_qs.aggregate(total=Sum("total_energy"))["total"] or Decimal("0")
        renewable_energy = energy_qs.aggregate(total=Sum("renewable_energy"))["total"] or Decimal("0")
        non_renewable_energy = energy_qs.aggregate(total=Sum("non_renewable_energy"))["total"] or Decimal("0")

        total_water = water_qs.aggregate(total=Sum("amount"))["total"] or Decimal("0")
        total_waste = waste_qs.aggregate(total=Sum("total_waste"))["total"] or Decimal("0")
        hazardous_waste = waste_qs.aggregate(total=Sum("hazardous"))["total"] or Decimal("0")
        non_hazardous_waste = waste_qs.aggregate(total=Sum("non_hazardous"))["total"] or Decimal("0")

        total_reduction = reduction_qs.aggregate(total=Sum("actual_reduction_tco2e"))["total"] or Decimal("0")
        total_offsets = credit_qs.aggregate(total=Sum("quantity_tco2e"))["total"] or Decimal("0")
        net_emissions = total_emissions - total_reduction - total_offsets

        scope_breakdown = emissions.values(
            "category__scope",
            "category__name",
            "scope3_category__name",
        ).annotate(
            total=Sum("total_co2e"),
            activity_count=Count("id"),
            quantity_total=Sum("quantity"),
        ).order_by("category__scope", "category__name", "scope3_category__name")

        facility_breakdown = emissions.values(
            "facility__name"
        ).annotate(
            total=Sum("total_co2e"),
            activity_count=Count("id")
        ).order_by("-total", "facility__name")

        category_breakdown = emissions.values(
            "category__name"
        ).annotate(
            total=Sum("total_co2e"),
            activity_count=Count("id")
        ).order_by("-total", "category__name")

        monthly_breakdown = self._build_monthly_breakdown(
            calculation_qs,
            org.name if org else "All Organizations",
            facility.name if facility else "All Facilities",
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "ESG Report"

        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        title_font = Font(size=18, bold=True, color="1F3864")
        section_font = Font(size=12, bold=True, color="FFFFFF")
        header_font = Font(size=11, bold=True, color="FFFFFF")
        normal_font = Font(size=11, color="000000")

        section_fill = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
        header_fill = PatternFill(fill_type="solid", start_color="4F81BD", end_color="4F81BD")
        light_fill = PatternFill(fill_type="solid", start_color="DCE6F1", end_color="DCE6F1")

        def apply_header(row_num, cols):
            """Style a worksheet row as a table header across the requested number of columns."""
            for col in range(1, cols + 1):
                cell = ws.cell(row=row_num, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

        def apply_section(row_num, start_col, end_col, label):
            """Create a merged section title row so each workbook block has a consistent visual heading."""
            ws.merge_cells(start_row=row_num, start_column=start_col, end_row=row_num, end_column=end_col)
            cell = ws.cell(row=row_num, column=start_col)
            cell.value = label
            cell.font = section_font
            cell.fill = section_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border

        def style_data_range(start_row, end_row, total_cols, number_cols=None, pct_cols=None):
            """Apply shared borders, fonts, and numeric formats to a rectangular block of report data."""
            number_cols = number_cols or []
            pct_cols = pct_cols or []
            for r in range(start_row, end_row + 1):
                for c in range(1, total_cols + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.font = normal_font
                    cell.border = border
                    if c in number_cols:
                        cell.number_format = '#,##0.000'
                    if c in pct_cols:
                        cell.number_format = '0.00%'

        for col, width in {
            "A": 20, "B": 24, "C": 18, "D": 18, "E": 18, "F": 18,
            "G": 18, "H": 18, "I": 18, "J": 18, "K": 18, "L": 18,
            "M": 18, "N": 18, "O": 18, "P": 18,
        }.items():
            ws.column_dimensions[col].width = width

        row = 1
        ws.merge_cells("A1:F2")
        ws["A1"] = "ESG Performance Report"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        # After merging A1:F2 for the title (rows 1-2), start content on row 3
        row = 3
        ws[f"A{row}"] = "Organization"
        ws[f"B{row}"] = org.name if org else "All Organizations"
        ws[f"D{row}"] = "Facility"
        ws[f"E{row}"] = facility.name if facility else "All Facilities"
        row += 1
        ws[f"A{row}"] = "Date Range"
        ws[f"B{row}"] = f"{start_date} to {end_date}"
        ws[f"D{row}"] = "Approved Records"
        ws[f"E{row}"] = emissions.count()
        row += 2

        apply_section(row, 1, 4, "EXECUTIVE SUMMARY")
        row += 1
        ws.append(["Metric", "Value", "Unit", "Comment"])
        apply_header(row, 4)
        row += 1

        summary_rows = [
            ["Total GHG Emissions", float(total_emissions), "tCO2e", "Approved activities within date range"],
            ["Scope 1 Emissions", float(s1_total), "tCO2e", ""],
            ["Scope 2 Emissions", float(s2_total), "tCO2e", ""],
            ["Scope 3 Emissions", float(s3_total), "tCO2e", ""],
            ["Total Energy", float(total_energy), "Unit from summaries", ""],
            ["Renewable Energy", float(renewable_energy), "Unit from summaries", ""],
            ["Non-Renewable Energy", float(non_renewable_energy), "Unit from summaries", ""],
            ["Total Water Usage", float(total_water), "Unit from summaries", ""],
            ["Total Waste", float(total_waste), "Unit from summaries", ""],
            ["Hazardous Waste", float(hazardous_waste), "Unit from summaries", ""],
            ["Non-Hazardous Waste", float(non_hazardous_waste), "Unit from summaries", ""],
            ["Emission Reductions", float(total_reduction), "tCO2e", ""],
            ["Offsets Retired", float(total_offsets), "tCO2e", ""],
            ["Net Emissions", float(net_emissions), "tCO2e", "Gross - reductions - offsets"],
        ]
        start_summary = row
        for item in summary_rows:
            ws.append(item)
            row += 1
        style_data_range(start_summary, row - 1, 4, number_cols=[2])

        row += 1
        apply_section(row, 1, 6, "COMPLETE SCOPE BREAKDOWN")
        row += 1
        ws.append(["Scope", "Category", "Scope 3 Category", "Total Quantity", "Activity Count", "Total tCO2e"])
        apply_header(row, 6)
        row += 1
        scope_start = row

        for item in scope_breakdown:
            ws.append([
                dict(EmissionScope.choices).get(item["category__scope"], item["category__scope"]),
                item["category__name"] or "N/A",
                item["scope3_category__name"] or "-",
                float(item["quantity_total"] or 0),
                item["activity_count"] or 0,
                float(item["total"] or 0),
            ])
            row += 1

        scope_end = row - 1
        if scope_end >= scope_start:
            style_data_range(scope_start, scope_end, 6, number_cols=[4, 6])

        row += 1
        apply_section(row, 1, 4, "CATEGORY BREAKDOWN")
        row += 1
        ws.append(["Category", "Total tCO2e", "Share of Total", "Activity Count"])
        apply_header(row, 4)
        row += 1
        category_start = row

        for item in category_breakdown:
            total = item["total"] or Decimal("0")
            share = (total / total_emissions) if total_emissions else Decimal("0")
            ws.append([
                item["category__name"] or "N/A",
                float(total),
                float(share),
                item["activity_count"] or 0,
            ])
            row += 1

        category_end = row - 1
        if category_end >= category_start:
            style_data_range(category_start, category_end, 4, number_cols=[2], pct_cols=[3])

        row += 1
        apply_section(row, 1, 3, "FACILITY BREAKDOWN")
        row += 1
        ws.append(["Facility", "Total tCO2e", "Activity Count"])
        apply_header(row, 3)
        row += 1
        facility_start = row

        for item in facility_breakdown:
            ws.append([
                item["facility__name"] or "N/A",
                float(item["total"] or 0),
                item["activity_count"] or 0,
            ])
            row += 1

        facility_end = row - 1
        if facility_end >= facility_start:
            style_data_range(facility_start, facility_end, 3, number_cols=[2])

        row += 1
        apply_section(row, 1, 4, "MONTHLY EMISSION TREND")
        row += 1
        ws.append(["Year", "Month", "Scope", "Total tCO2e"])
        apply_header(row, 4)
        row += 1
        monthly_start = row

        for item in monthly_breakdown:
            ws.append([
                item["year"],
                calendar.month_abbr[item["month"]],
                dict(EmissionScope.choices).get(item["scope"], item["scope"]),
                float(item["total_co2e"] or 0),
            ])
            row += 1

        monthly_end = row - 1
        if monthly_end >= monthly_start:
            style_data_range(monthly_start, monthly_end, 4, number_cols=[4])

        row += 1
        apply_section(row, 1, 5, "NEUTRALITY CLAIMS")
        row += 1
        ws.append(["Name", "Gross Emissions", "Reductions", "Offsets Retired", "Net Emissions"])
        apply_header(row, 5)
        row += 1
        neutrality_start = row

        for claim in neutrality_qs:
            ws.append([
                claim.name,
                float(claim.gross_emissions or 0),
                float(claim.reductions or 0),
                float(claim.offsets_retired or 0),
                float(claim.net_emissions or 0),
            ])
            row += 1

        neutrality_end = row - 1
        if neutrality_end >= neutrality_start:
            style_data_range(neutrality_start, neutrality_end, 5, number_cols=[2, 3, 4, 5])

        helper_col1 = 26
        helper_col2 = 27
        category_helper_col1 = 29
        category_helper_col2 = 30
        monthly_helper_col1 = 32
        monthly_helper_col2 = 33

        ws.cell(row=1, column=helper_col1, value="Scope")
        ws.cell(row=1, column=helper_col2, value="tCO2e")
        ws.cell(row=2, column=helper_col1, value="Scope 1")
        ws.cell(row=2, column=helper_col2, value=float(s1_total))
        ws.cell(row=3, column=helper_col1, value="Scope 2")
        ws.cell(row=3, column=helper_col2, value=float(s2_total))
        ws.cell(row=4, column=helper_col1, value="Scope 3")
        ws.cell(row=4, column=helper_col2, value=float(s3_total))

        positive_scope_totals = [value for value in (float(s1_total), float(s2_total), float(s3_total)) if value > 0]
        if len(positive_scope_totals) >= 2:
            pie = PieChart()
            pie.title = "Emissions by Scope"
            pie.height = 10
            pie.width = 12
            pie_data = Reference(ws, min_col=helper_col2, min_row=2, max_row=4)
            pie_labels = Reference(ws, min_col=helper_col1, min_row=2, max_row=4)
            pie.add_data(pie_data, titles_from_data=False)
            pie.set_categories(pie_labels)
            pie.legend.position = "r"
            pie_lbls = DataLabelList()
            pie_lbls.showPercent = False
            pie_lbls.showVal = False
            pie_lbls.showCatName = False
            pie_lbls.showLegendKey = False
            pie.dLbls = pie_lbls
            ws.add_chart(pie, "H4")
        else:
            ws.merge_cells("H4:K8")
            note_cell = ws["H4"]
            note_cell.value = "Emissions by Scope chart is hidden when fewer than two scopes have data in the selected range."
            note_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            note_cell.font = Font(size=11, italic=True, color="5B6470")
            note_cell.fill = light_fill
            note_cell.border = border

        if category_end >= category_start:
            ws.cell(row=1, column=category_helper_col1, value="Category")
            ws.cell(row=1, column=category_helper_col2, value="Total tCO2e")
            chart_categories = list(category_breakdown[:8])
            for index, item in enumerate(chart_categories, start=2):
                ws.cell(row=index, column=category_helper_col1, value=item["category__name"] or "N/A")
                ws.cell(row=index, column=category_helper_col2, value=float(item["total"] or 0))

            bar = BarChart()
            bar.title = "Category Emissions"
            bar.y_axis.title = "tCO2e"
            bar.x_axis.title = "Category"
            bar.height = 10
            bar.width = 16
            bar.legend = None
            chart_category_end_row = 1 + len(chart_categories)
            bar_data = Reference(ws, min_col=category_helper_col2, min_row=1, max_row=chart_category_end_row)
            bar_cats = Reference(ws, min_col=category_helper_col1, min_row=2, max_row=chart_category_end_row)
            bar.add_data(bar_data, titles_from_data=True)
            bar.set_categories(bar_cats)
            bar.style = 10
            ws.add_chart(bar, "H26")

        if monthly_end >= monthly_start:
            monthly_totals_for_chart = {}
            for item in monthly_breakdown:
                key = (item["year"], item["month"])
                monthly_totals_for_chart[key] = monthly_totals_for_chart.get(key, 0.0) + float(item["total_co2e"] or 0)

            monthly_chart_rows = [
                (f"{calendar.month_abbr[month]} {year}", total_value)
                for (year, month), total_value in monthly_totals_for_chart.items()
            ]

            ws.cell(row=1, column=monthly_helper_col1, value="Month")
            ws.cell(row=1, column=monthly_helper_col2, value="Total tCO2e")
            for index, (label, total_value) in enumerate(monthly_chart_rows, start=2):
                ws.cell(row=index, column=monthly_helper_col1, value=label)
                ws.cell(row=index, column=monthly_helper_col2, value=total_value)

            line = LineChart()
            line.title = "Monthly Emission Trend"
            line.y_axis.title = "tCO2e"
            line.x_axis.title = "Month"
            line.height = 10
            line.width = 16
            line.legend = None
            monthly_chart_end_row = 1 + len(monthly_chart_rows)
            line_data = Reference(ws, min_col=monthly_helper_col2, min_row=1, max_row=monthly_chart_end_row)
            line_cats = Reference(ws, min_col=monthly_helper_col1, min_row=2, max_row=monthly_chart_end_row)
            line.add_data(line_data, titles_from_data=True)
            line.set_categories(line_cats)
            line.style = 13
            ws.add_chart(line, "H48")

        raw_ws = wb.create_sheet("Raw Activities")
        raw_headers = [
            "Organization",
            "Facility",
            "Department",
            "Reporting Period",
            "Category",
            "Scope",
            "Scope 3 Category",
            "Activity Name",
            "Activity Date",
            "Quantity",
            "Unit",
            "Factor Code",
            "Scope 2 Method",
            "Data Quality",
            "Supplier Name",
            "Invoice Number",
            "Source Reference",
            "Total CO2e",
            "Status",
        ]
        raw_ws.append(raw_headers)

        for col in range(1, len(raw_headers) + 1):
            cell = raw_ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        raw_row = 2
        for obj in emissions.order_by("activity_date", "id"):
            raw_ws.append([
                obj.organization.name if obj.organization else "",
                obj.facility.name if obj.facility else "",
                obj.department.name if obj.department else "",
                obj.reporting_period.name if obj.reporting_period else "",
                obj.category.name if obj.category else "",
                obj.category.get_scope_display() if obj.category else "",
                obj.scope3_category.name if obj.scope3_category else "",
                obj.activity_name,
                obj.activity_date.strftime("%Y-%m-%d") if obj.activity_date else "",
                float(obj.quantity or 0),
                obj.unit.code if obj.unit else "",
                obj.factor_code,
                obj.get_scope2_method_display() if obj.scope2_method else "",
                obj.get_data_quality_display() if obj.data_quality else "",
                obj.supplier_name,
                obj.invoice_number,
                obj.source_reference,
                float(obj.total_co2e or 0),
                obj.get_status_display() if obj.status else "",
            ])
            raw_row += 1

        for col_letter, width in {
            "A": 20, "B": 20, "C": 18, "D": 18, "E": 20, "F": 12, "G": 22,
            "H": 24, "I": 14, "J": 14, "K": 12, "L": 16, "M": 18, "N": 16,
            "O": 20, "P": 18, "Q": 22, "R": 16, "S": 14
        }.items():
            raw_ws.column_dimensions[col_letter].width = width

        for r in range(2, raw_ws.max_row + 1):
            for c in range(1, len(raw_headers) + 1):
                raw_ws.cell(row=r, column=c).border = border
            raw_ws.cell(row=r, column=10).number_format = '#,##0.000'
            raw_ws.cell(row=r, column=18).number_format = '#,##0.000'

        monthly_ws = wb.create_sheet("Monthly Summary")
        monthly_headers = ["Organization", "Facility", "Year", "Month", "Scope", "Total CO2e"]
        monthly_ws.append(monthly_headers)

        for col in range(1, len(monthly_headers) + 1):
            cell = monthly_ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for item in monthly_breakdown:
            monthly_ws.append([
                item["organization"],
                item["facility"],
                item["year"],
                calendar.month_abbr[item["month"]],
                dict(EmissionScope.choices).get(item["scope"], item["scope"]),
                float(item["total_co2e"] or 0),
            ])

        for r in range(2, monthly_ws.max_row + 1):
            for c in range(1, len(monthly_headers) + 1):
                monthly_ws.cell(row=r, column=c).border = border
            monthly_ws.cell(row=r, column=6).number_format = '#,##0.000'

        for col_letter, width in {
            "A": 20, "B": 20, "C": 10, "D": 12, "E": 12, "F": 16
        }.items():
            monthly_ws.column_dimensions[col_letter].width = width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        file_name = f"esg_report_{start_date}_to_{end_date}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{file_name}"'
        return response

    def _build_monthly_breakdown(self, calculation_qs, organization_name, facility_name):
        """Build month rows from calculations only."""
        calculation_rows = (
            calculation_qs.values(
                "activity__organization__name",
                "activity__facility__name",
                "activity__category__scope",
            )
            .annotate(
                year_value=ExtractYear("activity__activity_date"),
                month_value=ExtractMonth("activity__activity_date"),
            )
            .annotate(total=Sum("co2e_amount"))
            .order_by("year_value", "month_value", "activity__category__scope")
        )

        if calculation_rows:
            return [
                {
                    "organization": row["activity__organization__name"] or organization_name,
                    "facility": row["activity__facility__name"] or facility_name,
                    "year": row["year_value"],
                    "month": row["month_value"],
                    "scope": row["activity__category__scope"],
                    "total_co2e": row["total"] or Decimal("0"),
                }
                for row in calculation_rows
            ]
        return []
